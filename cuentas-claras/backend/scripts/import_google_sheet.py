"""
Script de importación y sincronización desde Google Sheet de Cuentas Claras.
Requiere definir la variable de entorno GOOGLE_SHEET_URL con el enlace de exportación XLSX:
export GOOGLE_SHEET_URL="https://docs.google.com/spreadsheets/d/<SHEET_ID>/export?format=xlsx"

Procesa:
- AHORROS_26 (cuentas y saldos en moneda extranjera y ARS actualizados)
- ene26, feb26, mar26, abr26, may26, jun26, jul26, agos26, sep26 (transacciones mensuales completas)
"""

import os
import sys
from datetime import date, datetime
import urllib.request
import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))

from app import create_app, db
from app.models.category import Category
from app.models.household import Household, HouseholdMember
from app.models.saving import MonthlySavingAdjustment, Saving, SavingAccount
from app.models.transaction import Transaction
from app.models.user import User

SHEET_URL = os.environ.get("GOOGLE_SHEET_URL", "")
XLSX_PATH = "/tmp/cuentas_claras_updated.xlsx"

MONTH_MAPPING = [
    ("ene26", 1),
    ("feb26", 2),
    ("mar26", 3),
    ("abr26", 4),
    ("may26", 5),
    ("jun26", 6),
    ("jul26", 7),
    ("agos26", 8),
    ("sep26", 9),
]


def download_sheet():
    if os.path.exists(XLSX_PATH) and os.path.getsize(XLSX_PATH) > 1000:
        print(f"✅ Usando archivo ya descargado en {XLSX_PATH}")
        return
    print(f"📥 Descargando Google Sheet actualizado...")
    import subprocess
    subprocess.run(["curl", "-L", "-s", "-o", XLSX_PATH, SHEET_URL], check=True)
    print(f"✅ Guardado en {XLSX_PATH}")


def parse_num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("$", "").replace(" ", "")
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except (ValueError, TypeError):
        return None


def parse_date(v, default_year, default_month):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        s = v.strip()
        for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                pass
    return date(default_year, default_month, 15)


def get_or_create_category(name, cat_type, icon="📦"):
    cat = Category.query.filter(db.func.lower(Category.name) == name.lower()).first()
    if not cat:
        cat = Category(name=name, type=cat_type, icon=icon, sort_order=90)
        db.session.add(cat)
        db.session.flush()
        print(f"   ✨ Categoría creada: {name} ({cat_type})")
    return cat


def map_category_and_user(concept, section, default_user_i, user_t):
    cn = concept.lower().strip()

    # Determinar usuario:
    user = default_user_i
    if any(k in cn for k in ("tamara", "tamy", "sueldo t", "sac t", "tarj tam", "mant.cta.tamy", "mantcta tamy")):
        user = user_t
    elif cn.endswith(" t") or cn.endswith(" t."):
        user = user_t
    elif any(k in cn for k in ("ivan", "sueldo i", "sac i", "hs extra i", "extra i", "t visa bna i")):
        user = default_user_i

    # Categorización
    cat_name = "Varios"
    cat_type = "variable_expense"
    icon = "📦"

    if section == "income":
        cat_type = "income"
        if "sueldo" in cn or "sac" in cn:
            cat_name = "Salario"
            icon = "💼"
        elif "hs extra" in cn or "extra i" in cn:
            cat_name = "Hs Extra"
            icon = "⏱️"
        elif "extra" in cn:
            cat_name = "Adicionales"
            icon = "🎁"
        elif "auh" in cn or "afip" in cn:
            cat_name = "Otros Ingresos"
            icon = "💵"
        else:
            cat_name = "Otros Ingresos"
            icon = "💵"

    elif section == "fixed_expense":
        cat_type = "fixed_expense"
        if "cuota dpto" in cn or "alquiler" in cn:
            cat_name = "Alquiler / Cuota Dpto"
            icon = "🏠"
        elif "tarjeta" in cn or "tarj" in cn or "visa" in cn:
            cat_name = "Tarjetas de Crédito"
            icon = "💳"
        elif "expensas" in cn:
            cat_name = "Expensas"
            icon = "🏢"
        elif "metrogas" in cn or "gas" in cn:
            cat_name = "Gas"
            icon = "🔥"
        elif "edesur" in cn or "luz" in cn:
            cat_name = "Luz"
            icon = "💡"
        elif "abl" in cn:
            cat_name = "ABL / Impuestos"
            icon = "🏛️"
        elif "seguro" in cn:
            cat_name = "Seguros"
            icon = "🛡️"
        elif "mant" in cn:
            cat_name = "Mantenimiento Cuenta"
            icon = "🏦"
        else:
            cat_name = "Varios"
            icon = "📦"

    else:  # variable_expense (Gastos Extra)
        cat_type = "variable_expense"
        if any(w in cn for w in ("super", "súper", "comida", "verdul", "fruta", "carne", "carnicer", "pan", "panader", "arepa", "chipa", "pasta", "queso", "huevo", "feria", "semilla", "leche", "almuerzo", "cena", "desayuno", "chino", "dia", "coto", "carrefour", "pescaderia")):
            cat_name = "Comida / Supermercado"
            icon = "🛒"
        elif any(w in cn for w in ("delivery", "helado", "salida", "pizza", "cerveza", "birra", "piave", "bar", "cafe", "cafeter", "restauran", "mcdonald", "burger", "pedidosya", "rappi")):
            cat_name = "Delivery / Salidas"
            icon = "🍕"
        elif any(w in cn for w in ("uber", "taxi", "peaje", "nafta", "combustible", "estacionamiento", "sube", "colectivo", "pesjes")):
            cat_name = "Transporte"
            icon = "🚗"
        elif any(w in cn for w in ("farmacia", "farmacity", "remedio", "medic", "pañal", "pañales", "salud", "dentista", "optica")):
            cat_name = "Salud / Farmacia"
            icon = "💊"
        elif any(w in cn for w in ("ropa", "zapatill", "calzado", "indumentaria", "vestid", "pantalon", "remera", "interior", "clandestine")):
            cat_name = "Indumentaria"
            icon = "👕"
        elif any(w in cn for w in ("easy", "ferreter", "limpieza", "pileta", "flete", "lavadero", "llave", "copia", "mueble", "cuota oficina", "peluquer")):
            cat_name = "Hogar / Mantenimiento"
            icon = "🔨"
        elif any(w in cn for w in ("gemini", "spotify", "netflix", "cine", "boca", "futbol", "recital", "juego")):
            cat_name = "Entretenimiento"
            icon = "🎬"
        elif any(w in cn for w in ("mascota", "veterinar", "perro", "gato")):
            cat_name = "Mascotas"
            icon = "🐾"
        else:
            cat_name = "Varios"
            icon = "📦"

    category = get_or_create_category(cat_name, cat_type, icon)
    return category, user


def parse_sheet_dynamically(ws, sheet_name):
    """Detecta automáticamente la posición de ENTRADAS, GASTOS FIJOS y GASTOS EXTRA."""
    pos_ent, pos_fij, pos_ext = None, None, None
    for r in range(1, 30):
        for c in range(1, 15):
            val = str(ws.cell(r, c).value or '').strip().upper()
            if val in ('ENTRADAS', 'INGRESOS') and not pos_ent:
                pos_ent = (r, c)
            elif 'GASTOS FIJOS' in val and not pos_fij:
                pos_fij = (r, c)
            elif ('GASTOS EXTRA' in val or 'GASTOS VARIABLES' in val) and not pos_ext:
                pos_ext = (r, c)

    ent = []
    if pos_ent:
        r_ent, c_ent = pos_ent
        for r in range(r_ent + 2, r_ent + 30):
            c = ws.cell(r, c_ent).value
            a = ws.cell(r, c_ent + 1).value
            if c and 'TOTAL' in str(c).upper():
                break
            amt = parse_num(a)
            if c and amt and amt > 0:
                ent.append((str(c).strip(), amt))

    fij = []
    if pos_fij:
        r_fij, c_fij = pos_fij
        for r in range(r_fij + 2, r_fij + 35):
            c = ws.cell(r, c_fij).value
            raw_a = ws.cell(r, c_fij + 1).value
            if c and 'TOTAL' in str(c).upper():
                break
            if sheet_name == 'ene26' and c and 'Expensas 511' in str(c) and isinstance(raw_a, str):
                continue
            amt = parse_num(raw_a)
            if c and amt and amt > 0:
                fij.append((str(c).strip(), amt))

    ext = []
    if pos_ext:
        r_ext, c_ext = pos_ext
        for r in range(r_ext + 2, ws.max_row + 1):
            c = ws.cell(r, c_ext).value
            d = ws.cell(r, c_ext + 1).value
            a = ws.cell(r, c_ext + 2).value
            if c and 'TOTAL' in str(c).upper():
                break
            amt = parse_num(a)
            if (c or amt) and amt and amt > 0:
                ext.append((str(c or 'Gasto Extra').strip(), d, amt))

    return ent, fij, ext


def run_import():
    app = create_app()
    with app.app_context():
        print("🚀 Iniciando sincronización con Google Sheet original actualizado...")
        download_sheet()

        wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)

        # 1. Configurar Usuarios y Household
        ivan = User.query.filter(User.email.like("%ivan%")).first()
        if not ivan:
            ivan = User(
                name="Ivan Exequiel Carneiro",
                email="ivanexequielc@gmail.com",
                short_name="I",
                photo_url="https://lh3.googleusercontent.com/a/ACg8ocK1aFh2pY4L12ShiJ5SHyoXjkPJfKsdIIWEBwxNtNWbQudY41GyGg=s96-c",
            )
            ivan.set_password("123456")
            db.session.add(ivan)
            db.session.flush()
        else:
            ivan.short_name = "I"
            if not ivan.photo_url:
                ivan.photo_url = "https://lh3.googleusercontent.com/a/ACg8ocK1aFh2pY4L12ShiJ5SHyoXjkPJfKsdIIWEBwxNtNWbQudY41GyGg=s96-c"

        household = Household.query.first()
        if not household:
            household = Household(name="Casa", created_by=ivan.id)
            db.session.add(household)
            db.session.flush()

        tamara = User.query.filter(User.name.like("%tamara%") | User.email.like("%tamara%")).first()
        if not tamara:
            tamara = User(
                name="Tamara",
                email="tamara@cuentasclaras.app",
                short_name="T",
            )
            tamara.set_password("123456")
            db.session.add(tamara)
            db.session.flush()
            print("👤 Usuario Tamara verificado.")
        else:
            tamara.short_name = "T"

        # Asegurar membresías en Household
        for u in (ivan, tamara):
            mem = HouseholdMember.query.filter_by(household_id=household.id, user_id=u.id).first()
            if not mem:
                mem = HouseholdMember(household_id=household.id, user_id=u.id, role="admin")
                db.session.add(mem)

        db.session.commit()

        # 2. Limpiar transacciones y cuentas de ahorro anteriores
        print("🧹 Limpiando base para reemplazo fiel con datos actualizados...")
        Transaction.query.filter_by(household_id=household.id).delete()
        MonthlySavingAdjustment.query.filter_by(household_id=household.id).delete()
        Saving.query.filter_by(household_id=household.id).delete()
        SavingAccount.query.filter_by(household_id=household.id).delete()
        db.session.commit()

        # 3. Importar AHORROS_26 (Tenencias y cuentas actualizadas)
        print("\n💰 Importando cuentas de AHORROS_26...")
        ws_ahorros = wb["AHORROS_26"]
        accounts_data = [
            # (Row, Concept, Currency, Default balance, Default balance ARS)
            (3, "R$ FT", "BRL", 680.0, 202772.11),
            (4, "U$D NX", "USD", 1627.0, 2496060.75),
            (5, "U$D FT", "USD", 280.0, 429561.78),
            (6, "BullMarket", "ARS", 506000.0, 506000.0),
            (7, "Sobró 2025", "ARS", 240000.0, 240000.0),
        ]

        for r, default_name, default_curr, def_bal, def_bal_ars in accounts_data:
            c_name = ws_ahorros.cell(r, 1).value or default_name
            val_orig = parse_num(ws_ahorros.cell(r, 2).value) or def_bal
            val_ars = parse_num(ws_ahorros.cell(r, 3).value) or (val_orig if default_curr == "ARS" else def_bal_ars)

            acc = SavingAccount(
                name=str(c_name).strip(),
                currency=default_curr,
                household_id=household.id,
                owner_id=ivan.id,
            )
            db.session.add(acc)
            db.session.flush()

            sav = Saving(
                account_id=acc.id,
                household_id=household.id,
                month=9,
                year=2026,
                balance=val_orig,
                balance_ars=val_ars,
            )
            db.session.add(sav)
            print(f"   💵 Cuenta: {acc.name} ({acc.currency}) -> {val_orig:,.2f} (AR$ {val_ars:,.2f})")

        db.session.commit()

        # 4. Importar Transacciones Mes por Mes
        total_tx_count = 0
        for sheet_name, month_num in MONTH_MAPPING:
            ws = wb[sheet_name]
            print(f"\n📅 Procesando hoja '{sheet_name}' (Mes {month_num}/2026)...")

            ent, fij, ext = parse_sheet_dynamically(ws, sheet_name)

            # A. ENTRADAS
            for concept, amt in ent:
                cat, user = map_category_and_user(concept, "income", ivan, tamara)
                tx = Transaction(
                    user_id=user.id,
                    household_id=household.id,
                    category_id=cat.id,
                    date=date(2026, month_num, 5),
                    description=concept,
                    amount=amt,
                    currency="ARS",
                    type="income",
                    is_fixed=True if "sueldo" in concept.lower() else False,
                )
                db.session.add(tx)

            # B. GASTOS FIJOS
            for concept, amt in fij:
                cat, user = map_category_and_user(concept, "fixed_expense", ivan, tamara)
                tx = Transaction(
                    user_id=user.id,
                    household_id=household.id,
                    category_id=cat.id,
                    date=date(2026, month_num, 10),
                    description=concept,
                    amount=amt,
                    currency="ARS",
                    type="expense",
                    is_fixed=True,
                )
                db.session.add(tx)

            # C. GASTOS EXTRA
            for concept, d_val, amt in ext:
                row_date = parse_date(d_val, 2026, month_num)
                if row_date.year != 2026 or row_date.month != month_num:
                    row_date = date(2026, month_num, min(row_date.day, 28))

                cat, user = map_category_and_user(concept, "variable_expense", ivan, tamara)
                tx = Transaction(
                    user_id=user.id,
                    household_id=household.id,
                    category_id=cat.id,
                    date=row_date,
                    description=concept,
                    amount=amt,
                    currency="ARS",
                    type="expense",
                    is_fixed=False,
                )
                db.session.add(tx)

            month_tx_count = len(ent) + len(fij) + len(ext)
            total_tx_count += month_tx_count
            print(f"   ✅ {month_tx_count} transacciones ({len(ent)} entradas, {len(fij)} fijos, {len(ext)} extras)")

        db.session.commit()
        print(f"\n🎉 ¡SINCRONIZACIÓN EXITOSA! Total transacciones importadas: {total_tx_count}")


if __name__ == "__main__":
    run_import()
